#!/usr/bin/env python3
"""Extract and summarize published per-sample TMB; never invent a callable denominator.

Download SOURCE_URL separately, then pass --workbook PATH. Requires openpyxl only
for extraction. The checked-in CSV permits offline recomputation with the stdlib.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
from statistics import mean, median

SOURCE_URL = (
    "https://static-content.springer.com/esm/art%3A10.1038%2Fng.2529/"
    "MediaObjects/41588_2013_BFng2529_MOESM8_ESM.xlsx"
)
SOURCE_SHA256 = "d4f9b2da7302d803764ef16f3c00112cce1d86f623abe39044b71cea951fc948"
FIELDS = (
    "case_id",
    "mycn_amp_status",
    "total_exonic_mut_mb",
    "nonsilent_mut_mb",
    "covered_bases",
    "source_row",
)


def summarize(rows):
    """Keep MYCN status 9 (unknown) in the overall cohort, outside both subgroups."""
    result = []
    for group, status in (("all_high_risk", None), ("MYCNamp", "1"), ("MYCNnonamp", "0")):
        selected = [r for r in rows if status is None or str(r["mycn_amp_status"]) == status]
        for field in ("total_exonic_mut_mb", "nonsilent_mut_mb"):
            values = [float(r[field]) for r in selected]
            result.append(
                {
                    "group": group,
                    "measure": field,
                    "n": len(values),
                    "median": median(values),
                    "mean": mean(values),
                }
            )
    return result


def extract(workbook):
    from openpyxl import load_workbook

    if hashlib.sha256(Path(workbook).read_bytes()).hexdigest() != SOURCE_SHA256:
        raise ValueError("Workbook hash differs from the reviewed source")
    sheet = load_workbook(workbook, read_only=True, data_only=True)["Master_Table"]
    header = next(sheet.iter_rows(min_row=2, max_row=2, max_col=56, values_only=True))
    expected = {
        0: "Case USI",
        9: "MYCNamp",
        52: "Total per Mb",
        53: "Nonsilent per Mb",
        55: "Covered bases",
    }
    if any(header[i] != label for i, label in expected.items()):
        raise ValueError("Unexpected source columns")
    rows = []
    for index, row in enumerate(
        sheet.iter_rows(min_row=3, max_row=242, max_col=56, values_only=True), 3
    ):
        if not row[0] or row[9] not in {0, 1, 9}:
            raise ValueError(f"Unexpected sample or MYCN status at row {index}")
        rows.append(dict(zip(FIELDS, (row[0], row[9], row[52], row[53], row[55], index))))
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    rows = extract(args.workbook)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    with (args.output_dir / "tmb-pugh2013-sample-rates.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    result = {
        "reference": "PMID:23334666",
        "source_url": SOURCE_URL,
        "sha256": SOURCE_SHA256,
        "source_locator": "Supplementary Table 1; Master_Table!A3:BD242; MYCN J, rates BA/BB, coverage BD",
        "method": "Unweighted sample medians/means of published rates; no rounded-count conversions. Exclude five MYCN=9 cases only from subgroup summaries.",
        "results": summarize(rows),
    }
    (args.output_dir / "tmb-pugh2013-recomputed.json").write_text(
        json.dumps(result, indent=2) + "\n"
    )


if __name__ == "__main__":
    main()
